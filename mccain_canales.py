"""Listado de operaciones McCain por canal (Granjas / Food Service / Mayoristas).

Mismo costo que el tablero: costo crudo de Chess - acuerdos McCain - bonificación
por canal (10% Granjas en 81109 / 81064 / 81895). Solo lectura.

Uso:  python mccain_canales.py [--desde 2026-01] [--salida ruta.xlsx]
"""
import argparse, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.chdir(os.path.dirname(os.path.abspath(__file__)))
import numpy as np
import pandas as pd
import data_pipeline as dp

CANALES = ["GRANJAS", "FOOD SERVICE", "MAYORISTAS"]
SKUS_BONIF = set(a for r in dp.BONIF_CANAL_REGLAS for a in r["articulos"])

COLS = {
    "fechaComprobate": "Fecha", "mes": "Mes", "dsDocumento": "Comprobante",
    "nrodoc": "Nro", "idCliente": "Id cliente", "nombreCliente": "Cliente",
    "dsSubcanalMKT": "Subcanal", "dsVendedor": "Vendedor", "region": "Región",
    "idArticulo": "Id artículo", "dsArticulo": "Artículo",
    "sku_bonif_10": "SKU con bonif. 10% Granjas",
    "kilos": "Kilos", "kilos_cargo": "Kilos cobrados",
    "bultos_cargo": "Bultos cobrados",
    "precioUnitarioNeto": "Precio unit. neto", "subtotalNeto": "Facturación neta",
    "preciocomprant": "Precio compra Chess", "costo_bruto": "Costo bruto Chess",
    "ajuste_mccain": "Ajuste acuerdo McCain", "ajuste_bonif_canal": "Ajuste bonif. 10% canal",
    "costo_unitario": "Costo final", "cm": "CM", "cm_pct": "CM %",
}
PLATA = {"Precio unit. neto", "Facturación neta", "Precio compra Chess", "Costo bruto Chess",
         "Ajuste acuerdo McCain", "Ajuste bonif. 10% canal", "Costo final", "CM"}


def cargar():
    df = pd.read_parquet(dp.PARQUET_PATH)
    df = dp.agregar_marca_linea(df)
    df = dp.aplicar_acuerdos(df)
    df = dp.aplicar_bonificacion_canal(df)
    df["fechaComprobate"] = pd.to_datetime(df["fechaComprobate"])
    return df


def preparar(df, desde):
    m = df[df["proveedor"].astype(str).str.contains("MC CAIN", case=False)
           & df["dsCanalMkt"].isin(CANALES)].copy()
    if desde:
        m = m[m["fechaComprobate"] >= pd.Timestamp(desde + "-01")]
    m["mes"] = m["fechaComprobate"].dt.strftime("%Y-%m")
    m["costo_bruto"] = dp.costo_bruto(m)
    m["cm"] = m["subtotalNeto"] - m["costo_unitario"]
    m["cm_pct"] = np.where(m["subtotalNeto"] != 0, m["cm"] / m["subtotalNeto"] * 100, np.nan)
    m["sku_bonif_10"] = np.where(pd.to_numeric(m["idArticulo"]).isin(SKUS_BONIF), "SI", "NO")
    return m.sort_values(["fechaComprobate", "nrodoc"])


def resumen(m, claves):
    r = (m.groupby(claves).agg(**{
        "Líneas": ("subtotalNeto", "size"),
        "Clientes": ("idCliente", "nunique"),
        "Kilos": ("kilos", "sum"),
        "Facturación neta": ("subtotalNeto", "sum"),
        "Costo bruto Chess": ("costo_bruto", "sum"),
        "Ajuste acuerdo McCain": ("ajuste_mccain", "sum"),
        "Ajuste bonif. 10% canal": ("ajuste_bonif_canal", "sum"),
        "Costo final": ("costo_unitario", "sum"),
        "CM": ("cm", "sum"),
    }).reset_index())
    r["CM %"] = np.where(r["Facturación neta"] != 0, r["CM"] / r["Facturación neta"] * 100, np.nan)
    # CM % sin la bonificación de canal, para ver el efecto del 10%
    sin = r["CM"] - r["Ajuste bonif. 10% canal"]
    r["CM % sin bonif. 10%"] = np.where(r["Facturación neta"] != 0, sin / r["Facturación neta"] * 100, np.nan)
    return r.rename(columns={"dsCanalMkt": "Canal", "mes": "Mes", "idArticulo": "Id artículo",
                             "dsArticulo": "Artículo", "sku_bonif_10": "SKU con bonif. 10% Granjas"})


def formatear(ws, t):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for i, col in enumerate(t.columns, 1):
        L = get_column_letter(i)
        s = t[col]
        if pd.api.types.is_datetime64_any_dtype(s):
            fmt, w = "DD/MM/YYYY", 12
        else:
            w = min(max([len(str(col))] + s.astype(str).str.len().head(300).tolist()) + 2, 45)
            if not pd.api.types.is_numeric_dtype(s) or col.startswith("Id") or col == "Nro":
                fmt = None
            elif col in PLATA:
                fmt = '"$" #,##0;[Red]-"$" #,##0'
            elif col.startswith("CM %"):
                fmt = '0.0"%";[Red]-0.0"%"'
            elif col.startswith("Kilos") or col.startswith("Bultos"):
                fmt = "#,##0.0"
            else:
                fmt = "#,##0"
        ws.column_dimensions[L].width = max(w, 10)
        if fmt:
            for cell in ws[L][1:]:
                cell.number_format = fmt


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--desde", default=None, help="YYYY-MM")
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    df = cargar()
    m = preparar(df, a.desde)
    ult = df["fechaComprobate"].max()
    salida = a.salida or f"mccain_canales_{ult:%Y-%m-%d}.xlsx"

    hojas = {
        "Resumen canal-mes": resumen(m, ["dsCanalMkt", "mes"]),
        "Resumen por SKU": resumen(m, ["dsCanalMkt", "idArticulo", "dsArticulo", "sku_bonif_10"])
                            .sort_values(["Canal", "Facturación neta"], ascending=[True, False]),
    }
    for c in CANALES:
        hojas[c.title()] = m[m["dsCanalMkt"] == c][list(COLS)].rename(columns=COLS)

    with pd.ExcelWriter(salida, engine="openpyxl") as w:
        nota = pd.DataFrame({"Notas": [
            f"Operaciones McCain (proveedor 9 - MC CAIN ARGENTINA SA) en Granjas, Food Service y Mayoristas. Datos hasta el {ult:%d/%m/%Y}.",
            "Mismo costo que el tablero: costo bruto Chess - acuerdo McCain por cliente - bonificación 10% por canal.",
            f"La bonificación 10% aplica solo en GRANJAS y solo a los SKUs {', '.join(map(str, sorted(SKUS_BONIF)))} (columna 'SKU con bonif. 10% Granjas').",
            "Incluye facturas y notas de crédito (las NC restan). Costo bruto = precio compra × kilos (o bultos si es no pesable) cobrados; lo bonificado va a costo 0.",
            "'CM % sin bonif. 10%' muestra cómo quedaría el margen sin esa regla, para ver su efecto.",
        ]})
        nota.to_excel(w, sheet_name="Notas", index=False)
        w.sheets["Notas"].column_dimensions["A"].width = 130
        for n, t in hojas.items():
            t.to_excel(w, sheet_name=n[:31], index=False)
            formatear(w.sheets[n[:31]], t)
    print(salida)
    print(hojas["Resumen canal-mes"].to_string())


if __name__ == "__main__":
    main()
